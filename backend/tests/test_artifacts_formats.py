from __future__ import annotations

import io
import zipfile
from datetime import date
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from app.artifacts import (
    ArtifactError,
    CodeBlock,
    DocumentSpec,
    HeadingBlock,
    HtmlArtifactSpec,
    InsertBlockOperation,
    ListBlock,
    ParagraphBlock,
    ReplaceBlockOperation,
    TableBlock,
    WorkbookCell,
    WorkbookChart,
    WorkbookSheet,
    WorkbookSpec,
    export_artifact,
    export_docx,
    export_xlsx,
    import_docx,
    import_xlsx,
    parse_markdown,
    preview_artifact,
    recalculate_xlsx,
    render_document_html,
    render_markdown,
    revise_document,
    revise_workbook,
)
from app.artifacts.errors import UnsupportedFeatureError
from app.artifacts.safety import inspect_office_archive


def test_markdown_roundtrip_preserves_structure_and_paragraph_newlines():
    source = (
        "# Heading 2\n\nA line with $5.00 and a\nreal newline.\n\n- one\n- two\n\n| Name | Value |\n| --- | --- |\n| alpha | 3 |\n\n"
        + chr(96) * 3
        + "py\nx = 1\n"
        + chr(96) * 3
        + "\n"
    )
    spec = parse_markdown(source, title="Course notes")
    assert [block.type for block in spec.blocks] == [
        "heading",
        "paragraph",
        "list",
        "table",
        "code",
    ]
    assert spec.blocks[1].text == "A line with $5.00 and a\nreal newline."
    reread = parse_markdown(render_markdown(spec), title=spec.title)
    assert [block.model_dump(exclude={"id"}) for block in reread.blocks] == [
        block.model_dump(exclude={"id"}) for block in spec.blocks
    ]
    safe = render_document_html(
        DocumentSpec(title="t", blocks=[ParagraphBlock(id="p", text="<script>alert(1)</script>")])
    )
    assert "<script>" not in safe
    assert "&lt;script&gt;" in safe


def test_markdown_table_cells_roundtrip_pipes_and_newlines():
    spec = DocumentSpec(
        title="t",
        blocks=[
            TableBlock(
                id="t1",
                headers=["A", "B"],
                rows=[["a|b", "line1\nline2"], ["plain", "x"]],
            )
        ],
    )
    reread = parse_markdown(render_markdown(spec), title="t")
    assert [block.type for block in reread.blocks] == ["table"]
    assert reread.blocks[0].headers == ["A", "B"]
    assert reread.blocks[0].rows == [["a|b", "line1\nline2"], ["plain", "x"]]


def test_markdown_empty_heading_degrades_to_paragraph():
    spec = parse_markdown("# \n\nBody\n", title="t")
    assert [block.type for block in spec.blocks] == ["paragraph", "paragraph"]
    assert spec.blocks[0].text.strip() == "#"


def test_docx_heading_above_six_degrades_to_paragraph():
    from docx import Document as WordDocument
    from docx.enum.style import WD_STYLE_TYPE

    document = WordDocument()
    if "Heading 7" not in [style.name for style in document.styles]:
        document.styles.add_style("Heading 7", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Deep heading", style="Heading 7")
    document.add_paragraph("Body")
    buffer = io.BytesIO()
    document.save(buffer)
    imported = import_docx(buffer.getvalue())
    assert [block.type for block in imported.spec.blocks] == ["paragraph", "paragraph"]
    assert imported.spec.blocks[0].text == "Deep heading"


def test_docx_generation_import_preserves_supported_content_and_original():
    spec = DocumentSpec(
        title="Lab notes",
        blocks=[
            HeadingBlock(id="heading-a", level=2, text="Week 2"),
            ParagraphBlock(id="para-a", text="A line\nand a second **bold** line."),
            ListBlock(id="list-a", ordered=True, items=["first", "second"]),
            TableBlock(id="table-a", headers=["Item", "Amount"], rows=[["books", "$12.50"]]),
            CodeBlock(id="code-a", language="py", text="x = 1"),
        ],
    )
    spec.page_layout.width_inches = 8.27
    spec.page_layout.height_inches = 11.69
    spec.page_layout.margin_left_inches = 0.7
    original = export_docx(spec)
    imported = import_docx(original)
    assert imported.original_bytes == original
    assert len(imported.original_sha256) == 64
    assert imported.spec.title == "Lab notes"
    assert imported.spec.page_layout.width_inches == pytest.approx(8.27, abs=0.01)
    assert imported.spec.page_layout.margin_left_inches == pytest.approx(0.7, abs=0.01)
    assert [block.type for block in imported.spec.blocks] == [
        "heading",
        "paragraph",
        "list",
        "table",
        "code",
    ]
    assert imported.spec.blocks[1].text == "A line\nand a second **bold** line."
    assert imported.spec.blocks[2].items == ["first", "second"]
    assert imported.spec.blocks[3].rows == [["books", "$12.50"]]


def test_docx_import_refuses_embedded_images_before_rewrite():
    ordinary = export_docx(DocumentSpec(title="x", blocks=[ParagraphBlock(id="p", text="ok")]))
    payload = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(ordinary)) as source, zipfile.ZipFile(payload, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("word/media/image1.png", b"placeholder")
    original = payload.getvalue()
    with pytest.raises(UnsupportedFeatureError) as exc:
        import_docx(original)
    assert "embedded images" in exc.value.features
    assert original == payload.getvalue()


def test_xlsx_multisheet_formula_style_import_and_export():
    spec = WorkbookSpec(
        title="Lab data",
        sheets=[
            WorkbookSheet(
                id="s1",
                name="Measurements",
                cells=[
                    WorkbookCell(
                        id="a1",
                        row=1,
                        column=1,
                        value="Trial",
                        style={"bold": True, "fill_color": "DDEEFF"},
                    ),
                    WorkbookCell(id="b1", row=1, column=2, value="Mass", style={"bold": True}),
                    WorkbookCell(id="a2", row=2, column=1, value="A"),
                    WorkbookCell(
                        id="b2", row=2, column=2, value=3.5, style={"number_format": "$0.00"}
                    ),
                    WorkbookCell(id="b3", row=3, column=2, formula="=SUM(B2:B2)"),
                    WorkbookCell(
                        id="date",
                        row=4,
                        column=1,
                        value=date(2026, 10, 7),
                        style={"number_format": "yyyy-mm-dd"},
                    ),
                ],
                column_widths={"A": 24},
                freeze_panes="A2",
            ),
            WorkbookSheet(
                id="s2",
                name="Notes",
                cells=[WorkbookCell(id="note", row=1, column=1, value="See lab notebook")],
            ),
        ],
    )
    data = export_xlsx(spec)
    imported = import_xlsx(data)
    assert imported.original_bytes == data
    assert [sheet.name for sheet in imported.spec.sheets] == ["Measurements", "Notes"]
    cells = {cell.id: cell for cell in imported.spec.sheets[0].cells}
    assert cells["cell-001-3-2"].formula == "=SUM(B2:B2)"
    assert cells["cell-001-2-2"].style.number_format == "$0.00"
    assert imported.spec.sheets[0].freeze_panes == "A2"
    assert cells["cell-001-4-1"].value == date(2026, 10, 7)
    output = export_xlsx(imported.spec)
    workbook = load_workbook(io.BytesIO(output), data_only=False)
    assert workbook.sheetnames == ["Measurements", "Notes"]
    assert workbook["Measurements"]["B3"].value == "=SUM(B2:B2)"
    assert workbook["Measurements"].column_dimensions["A"].width == 24
    workbook.close()


def test_xlsx_multiline_wrap_and_row_heights_roundtrip_without_losing_alignment_safety():
    spec = WorkbookSpec(
        title="Wrapped notes",
        sheets=[
            WorkbookSheet(
                id="s",
                name="Notes",
                cells=[
                    WorkbookCell(id="auto", row=1, column=1, value="first\nsecond\nthird"),
                    WorkbookCell(
                        id="explicit-true", row=2, column=1, value="wrapped\ntext",
                        style={"wrap_text": True},
                    ),
                    WorkbookCell(
                        id="explicit-false", row=3, column=1, value="unwrapped\ntext",
                        style={"wrap_text": False},
                    ),
                    WorkbookCell(id="custom-height", row=4, column=1, value="one line"),
                    WorkbookCell(id="minimum-height", row=5, column=1, value="a\nb\nc"),
                ],
                row_heights={2: 42.0, 4: 28.5, 5: 20.0},
            )
        ],
    )
    data = export_xlsx(spec)
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    worksheet = workbook["Notes"]
    assert worksheet["A1"].alignment.wrap_text is True
    assert worksheet.row_dimensions[1].height == 45
    assert worksheet["A2"].alignment.wrap_text is True
    assert worksheet.row_dimensions[2].height == 42
    # OOXML omits an explicit false flag; import infers no-wrap for multiline text.
    assert worksheet["A3"].alignment.wrap_text is None
    assert worksheet.row_dimensions[4].height == 28.5
    assert worksheet.row_dimensions[5].height == 45
    workbook.close()

    imported = import_xlsx(data)
    cells = {cell.row: cell for cell in imported.spec.sheets[0].cells}
    assert cells[1].value == "first\nsecond\nthird"
    assert cells[1].style.wrap_text is True
    assert cells[3].style.wrap_text is False
    assert imported.spec.sheets[0].row_heights == {1: 45.0, 2: 42.0, 4: 28.5, 5: 45.0}
    reexport = load_workbook(io.BytesIO(export_xlsx(imported.spec)), data_only=False)
    assert reexport["Notes"]["A1"].alignment.wrap_text is True
    assert reexport["Notes"].row_dimensions[1].height == 45
    assert reexport["Notes"]["A3"].alignment.wrap_text is None
    reexport.close()

    centered = load_workbook(io.BytesIO(data))
    centered["Notes"]["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
    malformed_layout = io.BytesIO()
    centered.save(malformed_layout)
    centered.close()
    with pytest.raises(UnsupportedFeatureError) as exc:
        import_xlsx(malformed_layout.getvalue())
    assert "custom alignment or text wrapping" in exc.value.features


def test_workbook_row_height_bounds_and_wrap_text_type_are_validated():
    with pytest.raises(ValueError):
        WorkbookSheet(
            id="s", name="Notes", cells=[], row_heights={0: 20},
        )
    with pytest.raises(ValueError):
        WorkbookSheet(
            id="s", name="Notes", cells=[], row_heights={1: 410},
        )
    with pytest.raises(ValueError):
        WorkbookCell(
            id="bad-wrap", row=1, column=1, value="x", style={"wrap_text": "yes"},
        )
    with pytest.raises(ValueError, match="maximum row height"):
        export_xlsx(
            WorkbookSpec(
                title="Too many lines",
                sheets=[WorkbookSheet(
                    id="s", name="Notes", cells=[WorkbookCell(
                        id="too-tall", row=1, column=1, value="\n".join("line" for _ in range(28)),
                    )],
                )],
            )
        )


def test_chart_generation_is_supported_but_chart_import_is_refused():
    spec = WorkbookSpec(
        title="Chart sample",
        sheets=[
            WorkbookSheet(
                id="s",
                name="Data",
                cells=[
                    WorkbookCell(id="h1", row=1, column=1, value="Course"),
                    WorkbookCell(id="h2", row=1, column=2, value="Count"),
                    WorkbookCell(id="c1", row=2, column=1, value="Math"),
                    WorkbookCell(id="v1", row=2, column=2, value=5),
                    WorkbookCell(id="c2", row=3, column=1, value="Physics"),
                    WorkbookCell(id="v2", row=3, column=2, value=3),
                ],
            )
        ],
        charts=[
            WorkbookChart(
                id="chart-1",
                type="bar",
                title="Counts",
                sheet_id="s",
                data_range="B1:B3",
                categories_range="A2:A3",
            )
        ],
    )
    data = export_xlsx(spec)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert any(name.startswith("xl/charts/chart") for name in archive.namelist())
    with pytest.raises(UnsupportedFeatureError, match="cannot preserve safely") as exc:
        import_xlsx(data)
    assert "existing charts" in exc.value.features


def test_formula_preview_and_export_never_claim_calculated_values_without_isolated_runner():
    spec = WorkbookSpec(
        title="Calc",
        sheets=[
            WorkbookSheet(
                id="s",
                name="Sheet1",
                cells=[
                    WorkbookCell(id="x", row=1, column=1, value=1),
                    WorkbookCell(id="total", row=2, column=1, formula="=SUM(A1:A1)"),
                ],
            )
        ],
    )
    preview = preview_artifact(spec)
    assert preview["status"] == "needs_recalculation"
    formula = next(cell for cell in preview["sheets"][0]["cells"] if cell["formula"])
    assert formula["calculatedValue"] is None
    exported = export_artifact(spec, ".xlsx")
    assert exported.recalculation_status == "needs_recalculation"
    assert "not configured" in exported.message
    assert recalculate_xlsx(exported.content, spec).status == "needs_recalculation"


def test_revisions_target_explicit_ids_and_leave_parent_spec_unchanged():
    base = DocumentSpec(title="Draft", blocks=[ParagraphBlock(id="p-1", text="Before")])
    changed = revise_document(
        base,
        [
            ReplaceBlockOperation(
                type="replace_block", block_id="p-1", block=ParagraphBlock(id="p-1", text="After")
            ),
            InsertBlockOperation(
                type="insert_block",
                after_id="p-1",
                block=HeadingBlock(id="h-1", level=2, text="Next"),
            ),
        ],
    )
    assert base.blocks[0].text == "Before"
    assert [block.id for block in changed.blocks] == ["p-1", "h-1"]
    with pytest.raises(ArtifactError, match="No block has ID"):
        revise_document(
            base,
            [
                ReplaceBlockOperation(
                    type="replace_block",
                    block_id="missing",
                    block=ParagraphBlock(id="missing", text="x"),
                )
            ],
        )
    workbook = WorkbookSpec(
        title="Sheet",
        sheets=[
            WorkbookSheet(
                id="s", name="Sheet1", cells=[WorkbookCell(id="a", row=1, column=1, value=1)]
            )
        ],
    )
    updated = revise_workbook(
        workbook,
        [
            {
                "type": "set_cell",
                "sheet_id": "s",
                "cell_id": "a",
                "cell": {"id": "a", "row": 1, "column": 1, "value": 2},
            },
            {
                "type": "set_cell",
                "sheet_id": "s",
                "cell_id": "b",
                "cell": {"id": "b", "row": 2, "column": 1, "formula": "=A1*2"},
            },
        ],
    )
    assert workbook.sheets[0].cells[0].value == 1
    assert [cell.id for cell in updated.sheets[0].cells] == ["a", "b"]


def test_office_archive_rejects_path_traversal_and_macro_parts():
    bad_path = io.BytesIO()
    with zipfile.ZipFile(bad_path, "w") as archive:
        archive.writestr("../evil.xml", "<x/>")
    with pytest.raises(ValueError, match="unsafe path"):
        inspect_office_archive(bad_path.getvalue(), kind="docx")
    macro = io.BytesIO()
    with zipfile.ZipFile(macro, "w") as archive:
        archive.writestr("xl/vbaProject.bin", b"MZ")
        archive.writestr("xl/workbook.xml", "<x/>")
    with pytest.raises(UnsupportedFeatureError) as exc:
        inspect_office_archive(macro.getvalue(), kind="xlsx")
    assert "macro-enabled content" in exc.value.features


def test_html_artifact_is_explicitly_marked_for_sandbox_preview():
    spec = HtmlArtifactSpec(
        title="Interactive lesson",
        html="<!doctype html><script>document.body.textContent='ok'</script>",
    )
    preview = preview_artifact(spec)
    assert preview["sandboxRequired"] is True
    exported = export_artifact(spec, "html")
    from html.parser import HTMLParser
    class Frames(HTMLParser):
        def handle_starttag(self, tag, attrs):
            if tag == "iframe":
                self.frame = dict(attrs)
    parser = Frames()
    parser.feed(exported.content.decode("utf-8"))
    assert parser.frame["sandbox"] == ""
    assert "document.body.textContent" not in parser.frame["srcdoc"]
    assert "script-src 'none'" in parser.frame["srcdoc"]
    assert "connect-src 'none'" in parser.frame["srcdoc"]
    assert "frame-src 'self'" in exported.content.decode("utf-8")
    assert exported.filename == "interactive-lesson.html"


@pytest.mark.parametrize(
    ("returned_formula", "changed_input"),
    [("SUM(A1:A1)", False), ("SUM(A1:A1)+1", False), ("SUM(A1:A1)", True)],
)
def test_recalculation_preserves_formulas_and_inputs_before_accepting_cached_values(
    returned_formula, changed_input,
):
    spec = WorkbookSpec(
        title="Calc",
        sheets=[
            WorkbookSheet(
                id="s",
                name="Sheet1",
                cells=[
                    WorkbookCell(id="x", row=1, column=1, value=1),
                    WorkbookCell(id="total", row=2, column=1, formula="=SUM(A1:A1)"),
                ],
            )
        ],
    )
    raw = export_xlsx(spec)
    with zipfile.ZipFile(io.BytesIO(raw)) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    sheet_xml = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    assert "<f>SUM(A1:A1)</f><v></v>" in sheet_xml
    recalculated_xml = sheet_xml.replace(
        "<f>SUM(A1:A1)</f><v></v>", f"<f>{returned_formula}</f><v>1</v>"
    )
    if changed_input:
        root = ElementTree.fromstring(recalculated_xml)
        namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        input_cell = root.find(f".//{namespace}c[@r='A1']/{namespace}v")
        assert input_cell is not None and input_cell.text == "1"
        input_cell.text = "9"
        recalculated_xml = ElementTree.tostring(root, encoding="unicode")
    parts["xl/worksheets/sheet1.xml"] = recalculated_xml.encode("utf-8")
    cached = io.BytesIO()
    with zipfile.ZipFile(cached, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)

    class IsolatedRunner:
        def recalculate_xlsx(self, xlsx_bytes: bytes, *, timeout_seconds: int) -> bytes:
            assert xlsx_bytes == raw
            assert timeout_seconds == 60
            return cached.getvalue()

    result = recalculate_xlsx(raw, spec, runner=IsolatedRunner())
    if returned_formula == "SUM(A1:A1)" and not changed_input:
        assert result.status == "verified"
        assert result.xlsx_bytes == cached.getvalue()
    else:
        assert result.status == "needs_recalculation"
        assert result.xlsx_bytes == raw
        changed_address = "Sheet1!A1" if changed_input else "Sheet1!A2"
        assert "changed formulas or inputs at: " + changed_address in result.message

    class FailingRunner:
        def recalculate_xlsx(self, xlsx_bytes: bytes, *, timeout_seconds: int) -> bytes:
            raise RuntimeError("sandbox unavailable")

    failed = recalculate_xlsx(raw, spec, runner=FailingRunner())
    assert failed.status == "needs_recalculation"
    assert failed.xlsx_bytes == raw
    assert "sandbox unavailable" in failed.message
