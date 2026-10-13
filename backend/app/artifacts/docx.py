"""DOCX import/export for the explicitly supported text structures."""

from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass

from defusedxml import ElementTree as SafeET
from docx import Document as WordDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .errors import UnsupportedFeatureError
from .models import (
    CodeBlock,
    DocumentBlock,
    DocumentSpec,
    HeadingBlock,
    ListBlock,
    PageLayoutSpec,
    ParagraphBlock,
    QuoteBlock,
    TableBlock,
)
from .safety import inspect_office_archive, read_bounded_entry


@dataclass(frozen=True)
class ImportResult:
    spec: DocumentSpec
    original_bytes: bytes
    original_sha256: str
    warnings: tuple[str, ...] = ()


def _inline_to_markdown(paragraph: Paragraph) -> str:
    parts = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        if run.underline or run.font.strike or run.font.all_caps:
            raise UnsupportedFeatureError(
                ["underlined, struck-through, or all-caps text formatting"]
            )
        if run.bold:
            text = "**" + text + "**"
        if run.italic:
            text = "*" + text + "*"
        if run.font.name in {"Consolas", "Courier New", "Courier"}:
            text = chr(96) + text + chr(96)
        parts.append(text)
    return "".join(parts)


_UNSUPPORTED_XML_FEATURES = (
    (("hyperlink",), "hyperlinks"),
    (("drawing",), "drawing objects"),
    (("altChunk",), "embedded documents"),
    (("gridSpan", "vMerge"), "merged table cells"),
    (("customXml",), "custom XML content"),
)


def _reject_unsupported_xml(main_root) -> None:
    local_names = {element.tag.rsplit("}", 1)[-1] for element in main_root.iter()}
    unsupported = {
        label for tags, label in _UNSUPPORTED_XML_FEATURES if local_names.intersection(tags)
    }
    if unsupported:
        raise UnsupportedFeatureError(sorted(unsupported))


def _reject_unsupported_layout(document) -> None:
    if len(document.sections) != 1:
        raise UnsupportedFeatureError(["multiple document sections"])
    for section in document.sections:
        header_content = [p.text for p in section.header.paragraphs if p.text.strip()]
        footer_content = [p.text for p in section.footer.paragraphs if p.text.strip()]
        if header_content or footer_content or section.header.tables or section.footer.tables:
            raise UnsupportedFeatureError(["non-empty headers or footers"])
    for table in document.tables:
        if any(cell.tables for row in table.rows for cell in row.cells):
            raise UnsupportedFeatureError(["nested tables"])
        for row in table.rows:
            if len(row.cells) != len(table.columns):
                raise UnsupportedFeatureError(["merged or irregular table cells"])


class _BlockCollector:
    """Accumulates imported blocks, numbering them and grouping consecutive list items."""

    def __init__(self):
        self.blocks: list[DocumentBlock] = []
        self._items: list[str] = []
        self._ordered = False
        self._kind: str | None = None

    def add(self, block_type, **values) -> None:
        self.blocks.append(block_type(id=f"block-{len(self.blocks) + 1:05d}", **values))

    def list_item(self, kind: str, text: str) -> None:
        if self._kind != kind:
            self.flush_list()
            self._kind = kind
            self._ordered = kind == "ordered"
        self._items.append(text)

    def flush_list(self) -> None:
        if self._items:
            self.add(ListBlock, ordered=self._ordered, items=self._items.copy())
            self._items = []
        self._kind = None


def _import_paragraph(paragraph, collector: _BlockCollector) -> str | None:
    """Add a paragraph's block; returns a document title if it is the leading Title paragraph."""
    style_name = paragraph.style.name if paragraph.style else "Normal"
    if style_name.startswith(("List Bullet", "List Number")):
        kind = "ordered" if style_name.startswith("List Number") else "unordered"
        collector.list_item(kind, _inline_to_markdown(paragraph))
        return None
    collector.flush_list()
    text = _inline_to_markdown(paragraph)
    if paragraph.runs and all(
        run.font.name == "Consolas" for run in paragraph.runs if run.text.strip()
    ):
        collector.add(CodeBlock, text=paragraph.text)
    elif style_name.startswith("Heading ") and style_name[8:].isdigit():
        level = int(style_name[8:])
        if 1 <= level <= 6:
            collector.add(HeadingBlock, level=level, text=text or " ")
        else:
            collector.add(ParagraphBlock, text=text)
    elif style_name == "Title":
        # Title paragraph style is represented by the spec title when it is first.
        if not collector.blocks and text.strip():
            return text.strip()
        collector.add(ParagraphBlock, text=text)
    elif style_name == "Quote":
        collector.add(QuoteBlock, text=text)
    else:
        collector.add(ParagraphBlock, text=text)
    return None


def import_docx(data: bytes, *, title: str | None = None) -> ImportResult:
    inspect_office_archive(data, kind="docx")
    main_xml = read_bounded_entry(data, "word/document.xml")
    try:
        main_root = SafeET.fromstring(main_xml)
    except Exception as exc:
        raise ValueError("DOCX document XML is unsafe or invalid") from exc
    _reject_unsupported_xml(main_root)
    try:
        document = WordDocument(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("DOCX content could not be read") from exc
    _reject_unsupported_layout(document)

    try:
        section = document.sections[0]
        # A crafted file may omit w:pgSz or margins, leaving None; keep the
        # spec defaults instead of raising AttributeError as a server error.
        measures = {
            "width_inches": section.page_width,
            "height_inches": section.page_height,
            "margin_top_inches": section.top_margin,
            "margin_bottom_inches": section.bottom_margin,
            "margin_left_inches": section.left_margin,
            "margin_right_inches": section.right_margin,
        }
        page_layout = PageLayoutSpec(
            **{key: value.inches for key, value in measures.items() if value is not None}
        )
    except (IndexError, AttributeError, TypeError) as exc:
        raise ValueError("DOCX page layout could not be read") from exc
    document_title = title
    collector = _BlockCollector()
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            leading_title = _import_paragraph(Paragraph(element, document), collector)
            document_title = leading_title or document_title
        elif element.tag == qn("w:tbl"):
            collector.flush_list()
            table = Table(element, document)
            rows = [
                ["\n".join(_inline_to_markdown(p) for p in cell.paragraphs) for cell in row.cells]
                for row in table.rows
            ]
            if rows:
                collector.add(TableBlock, headers=rows[0], rows=rows[1:])
    collector.flush_list()
    return ImportResult(
        spec=DocumentSpec(
            title=document_title or "Imported document", blocks=collector.blocks, page_layout=page_layout
        ),
        original_bytes=data,
        original_sha256=hashlib.sha256(data).hexdigest(),
    )


def _add_inline(paragraph, text: str) -> None:
    """Render the small supported inline subset emitted by Markdown/DOCX import."""
    pattern = re.compile(
        r"(\*\*.+?\*\*|(?<!\*)\*[^*]+\*(?!\*)|" + chr(96) + r"[^" + chr(96) + r"]+" + chr(96) + r")"
    )
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def export_docx(spec: DocumentSpec) -> bytes:
    document = WordDocument()
    for name in ("Title", "Subtitle", *(f"Heading {level}" for level in range(1, 10))):
        style = document.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style.element.pPr
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    document.core_properties.title = spec.title
    section = document.sections[0]
    section.page_width = Inches(spec.page_layout.width_inches)
    section.page_height = Inches(spec.page_layout.height_inches)
    section.top_margin = Inches(spec.page_layout.margin_top_inches)
    section.bottom_margin = Inches(spec.page_layout.margin_bottom_inches)
    section.left_margin = Inches(spec.page_layout.margin_left_inches)
    section.right_margin = Inches(spec.page_layout.margin_right_inches)
    title = document.add_paragraph(spec.title, style="Title")
    properties = title._p.pPr
    if properties is not None:
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
    for block in spec.blocks:
        if isinstance(block, HeadingBlock):
            _add_inline(document.add_paragraph(style=f"Heading {block.level}"), block.text)
        elif isinstance(block, ParagraphBlock):
            para = document.add_paragraph()
            _add_inline(para, block.text)
        elif isinstance(block, ListBlock):
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                _add_inline(document.add_paragraph(style=style), item)
        elif isinstance(block, QuoteBlock):
            _add_inline(document.add_paragraph(style="Quote"), block.text)
        elif isinstance(block, CodeBlock):
            para = document.add_paragraph()
            run = para.add_run(block.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif isinstance(block, TableBlock):
            table = document.add_table(rows=1, cols=len(block.headers))
            table.style = "Light Shading Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(table.rows[0].cells, block.headers, strict=True):
                _add_inline(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    _add_inline(cell.paragraphs[0], text)
    if spec.evidence:
        from app.artifacts.evidence import source_legend

        document.add_paragraph("Sources", style="Heading 1")
        for element, source, locator, chunk_id in source_legend(spec):
            document.add_paragraph(f"{element}: {source} · {locator} ({chunk_id})")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
