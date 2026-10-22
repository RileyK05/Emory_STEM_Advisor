"""XLSX import/export for bounded, explicit workbook structures."""

from __future__ import annotations

import hashlib
import io
import re
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from typing import Protocol

from defusedxml import ElementTree as SafeET
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter, range_boundaries

from .errors import ArtifactLimitError, UnsupportedFeatureError
from .models import CellStyle, WorkbookCell, WorkbookSheet, WorkbookSpec
from .safety import inspect_office_archive, read_bounded_entry

MAX_WORKBOOK_CELLS = 100_000
_FORMULA_BLOCKLIST = re.compile(
    # The bare pipe covers legacy DDE formulas (`=cmd|'/c calc'!A1`), which the
    # external-reference bracket rule alone does not catch.
    r"\b(?:WEBSERVICE|FILTERXML|RTD|CALL|REGISTER|EXEC|HYPERLINK)\s*\(|\[[^]]+\]|\|",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class WorkbookImportResult:
    spec: WorkbookSpec
    original_bytes: bytes
    original_sha256: str
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class RecalculationResult:
    status: str
    xlsx_bytes: bytes
    message: str | None = None


class IsolatedLibreOfficeRunner(Protocol):
    """Application boundary implemented by the separately isolated sandbox host."""

    def recalculate_xlsx(self, xlsx_bytes: bytes, *, timeout_seconds: int) -> bytes: ...


def _preflight_sheet_xml(data: bytes, names: frozenset[str]) -> None:
    total_cells = 0
    max_row_seen = 1
    max_column_seen = 1
    for name in sorted(
        n for n in names if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")
    ):
        xml = read_bounded_entry(data, name)
        try:
            root = SafeET.fromstring(xml)
        except Exception as exc:
            raise ValueError("XLSX worksheet XML is invalid") from exc
        cell_refs = [
            element.attrib.get("r")
            for element in root.iter()
            if element.tag.rsplit("}", 1)[-1] == "c"
        ]
        count = len(cell_refs)
        for ref in cell_refs:
            if ref:
                column, row = range_boundaries(ref + ":" + ref)[:2]
                # range_boundaries returns min_col, min_row, max_col, max_row.
                max_column_seen = max(max_column_seen, column)
                max_row_seen = max(max_row_seen, row)
        total_cells += count
        if max_row_seen * max_column_seen > 2_000_000:
            raise ArtifactLimitError("Workbook used-cell rectangle is too large for safe import.")
        if total_cells > MAX_WORKBOOK_CELLS:
            raise ArtifactLimitError("Workbook contains too many populated or styled cells.")
        unsupported_tags = {"dataValidations", "conditionalFormatting", "extLst", "legacyDrawing"}
        found = sorted(
            {
                element.tag.rsplit("}", 1)[-1]
                for element in root.iter()
                if element.tag.rsplit("}", 1)[-1] in unsupported_tags
            }
        )
        if found:
            raise UnsupportedFeatureError(["worksheet features: " + ", ".join(found)])


def _rgb(color) -> str | None:
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str):
        value = color.rgb[-6:]
        return value.upper()
    return None


def _check_cell_features(cell) -> None:
    if cell.comment is not None:
        raise UnsupportedFeatureError(["cell comments"])
    if cell.hyperlink is not None:
        raise UnsupportedFeatureError(["hyperlinks"])
    if cell.data_type == "f" and _FORMULA_BLOCKLIST.search(str(cell.value)):
        raise UnsupportedFeatureError(["external or execution-capable formulas"])
    if cell.has_style:
        border = copy(cell.border)
        if any(
            getattr(border, side).style or getattr(border, side).color
            for side in ("left", "right", "top", "bottom", "diagonal")
        ):
            raise UnsupportedFeatureError(["cell borders"])
        alignment = copy(cell.alignment)
        if (
            alignment.horizontal not in (None, "general")
            or alignment.vertical not in (None, "bottom")
            or alignment.textRotation
            or alignment.shrinkToFit
            or alignment.indent
        ):
            raise UnsupportedFeatureError(["custom alignment or text wrapping"])
        protection = copy(cell.protection)
        if not protection.locked or protection.hidden:
            raise UnsupportedFeatureError(["cell protection"])
        if cell.font.underline or cell.font.strike or cell.font.vertAlign:
            raise UnsupportedFeatureError(["unsupported font effects"])
        if cell.font.color is not None and _rgb(cell.font.color) is None:
            raise UnsupportedFeatureError(["theme or indexed font colors"])
        if cell.fill.fill_type not in (None, "solid"):
            raise UnsupportedFeatureError(["pattern fills"])
        if cell.fill.fill_type == "solid" and cell.fill.fgColor.type != "rgb":
            raise UnsupportedFeatureError(["theme or indexed cell fills"])


def _import_value(cell):
    """A plain cell's value, with midnight datetimes in date formats narrowed to dates."""
    value = cell.value
    if (
        isinstance(value, datetime)
        and value.time() == datetime.min.time()
        and is_date_format(cell.number_format)
        and not any(token in cell.number_format.casefold() for token in ("h", "s", "am/pm"))
    ):
        value = value.date()
    if value is not None and not isinstance(value, (str, int, float, bool, date, datetime)):
        raise UnsupportedFeatureError(["unsupported cell value type"])
    return value


def _import_cell(cell, sheet_index: int) -> WorkbookCell:
    _check_cell_features(cell)
    if cell.data_type == "f":
        formula = str(cell.value)
        content = {"formula": formula if formula.startswith("=") else "=" + formula}
    else:
        content = {"value": _import_value(cell)}
    return WorkbookCell(
        id=f"cell-{sheet_index:03d}-{cell.row}-{cell.column}",
        row=cell.row,
        column=cell.column,
        style=_cell_style(cell),
        **content,
    )


def _import_sheet(worksheet, sheet_index: int) -> WorkbookSheet:
    if worksheet.merged_cells.ranges:
        raise UnsupportedFeatureError(["merged cells"])
    # Iterate only the cells declared in the sheet XML. Walking the used
    # rectangle with iter_rows() makes openpyxl materialize a Cell for every
    # empty coordinate, so a sparse sheet (one cell at the far edge) can be a
    # memory bomb. `_cells` holds exactly the declared cells the preflight
    # already counted.
    cells = [
        _import_cell(cell, sheet_index)
        for cell in sorted(worksheet._cells.values(), key=lambda c: (c.row, c.column))
        if cell.value is not None or cell.has_style
    ]
    freeze = worksheet.freeze_panes
    return WorkbookSheet(
        id=f"sheet-{sheet_index:03d}",
        name=worksheet.title,
        cells=cells,
        column_widths={
            key: float(dim.width)
            for key, dim in worksheet.column_dimensions.items()
            if dim.width is not None and dim.width != 13
        },
        row_heights={
            int(key): float(dim.height)
            for key, dim in worksheet.row_dimensions.items()
            if dim.height is not None
        },
        freeze_panes=freeze if isinstance(freeze, str) else (freeze.coordinate if freeze else None),
    )


def import_xlsx(data: bytes, *, title: str = "Imported workbook") -> WorkbookImportResult:
    archive = inspect_office_archive(data, kind="xlsx")
    _preflight_sheet_xml(data, archive.names)
    try:
        workbook = load_workbook(
            io.BytesIO(data), data_only=False, read_only=False, keep_links=False
        )
    except Exception as exc:
        raise ValueError("XLSX content could not be read") from exc
    if len(workbook.worksheets) > 100:
        raise ArtifactLimitError("Workbook contains more than 100 worksheets.")

    sheets: list[WorkbookSheet] = []
    cell_count = 0
    for sheet_index, worksheet in enumerate(workbook.worksheets, 1):
        sheet = _import_sheet(worksheet, sheet_index)
        cell_count += len(sheet.cells)
        if cell_count > MAX_WORKBOOK_CELLS:
            raise ArtifactLimitError("Workbook contains too many populated or styled cells.")
        sheets.append(sheet)
    workbook.close()
    return WorkbookImportResult(
        spec=WorkbookSpec(title=title, sheets=sheets),
        original_bytes=data,
        original_sha256=hashlib.sha256(data).hexdigest(),
    )


def _cell_style(cell) -> CellStyle:
    color = _rgb(cell.font.color)
    fill = _rgb(cell.fill.fgColor) if cell.fill.fill_type == "solid" else None
    return CellStyle(
        bold=bool(cell.font.bold),
        italic=bool(cell.font.italic),
        font_color=color,
        fill_color=fill,
        number_format=cell.number_format if cell.number_format != "General" else None,
        wrap_text=(
            False
            if cell.alignment.wrapText is None
            and isinstance(cell.value, str)
            and re.search(r"\r|\n", cell.value)
            else cell.alignment.wrapText
        ),
    )


def _apply_style(cell, style: CellStyle) -> None:
    cell.font = Font(bold=style.bold, italic=style.italic, color=style.font_color)
    if style.fill_color:
        cell.fill = PatternFill(fill_type="solid", fgColor=style.fill_color)
    if style.number_format:
        cell.number_format = style.number_format
    cell.alignment = Alignment(wrap_text=style.wrap_text)


def _validate_chart_range(
    value: str, *, row_count: int, column_count: int
) -> tuple[int, int, int, int]:
    min_col, min_row, max_col, max_row = range_boundaries(value)
    if (
        min_row > row_count
        or max_row > row_count
        or min_col > column_count
        or max_col > column_count
    ):
        raise ValueError("chart range points outside populated workbook cells")
    if min_row == max_row and min_col == max_col:
        raise ValueError("chart range must include more than one cell")
    return min_col, min_row, max_col, max_row


def _write_sheet(workbook: Workbook, sheet_spec: WorkbookSheet):
    worksheet = workbook.create_sheet(sheet_spec.name)
    for cell_spec in sheet_spec.cells:
        cell = worksheet.cell(row=cell_spec.row, column=cell_spec.column)
        if cell_spec.formula and _FORMULA_BLOCKLIST.search(cell_spec.formula):
            raise UnsupportedFeatureError(["external or execution-capable formulas"])
        cell.value = cell_spec.formula if cell_spec.formula is not None else cell_spec.value
        if cell_spec.formula is None and isinstance(cell_spec.value, str):
            cell.data_type = "s"  # A literal beginning with = must remain inert text.
        _apply_style(cell, cell_spec.style)
        if (
            cell_spec.formula is None
            and isinstance(cell_spec.value, str)
            and re.search(r"\r|\n", cell_spec.value)
        ):
            line_count = len(re.split(r"\r\n|\r|\n", cell_spec.value))
            if cell_spec.style.wrap_text is None:
                cell.alignment = Alignment(wrap_text=True)
            if cell_spec.style.wrap_text is not False:
                minimum_height = 15 * line_count
                if minimum_height > 409:
                    raise ValueError("multiline cell exceeds Excel's maximum row height")
                requested_height = sheet_spec.row_heights.get(cell_spec.row, 0)
                worksheet.row_dimensions[cell_spec.row].height = max(
                    requested_height, minimum_height
                )
    for row, height in sheet_spec.row_heights.items():
        minimum = worksheet.row_dimensions[row].height or 0
        worksheet.row_dimensions[row].height = max(height, minimum)
    for column, width in sheet_spec.column_widths.items():
        if not re.fullmatch(r"[A-Z]{1,3}", column) or not 0 < width <= 255:
            raise ValueError("column widths require A1-style columns and values in (0, 255]")
        worksheet.column_dimensions[column].width = width
    if sheet_spec.freeze_panes:
        worksheet.freeze_panes = sheet_spec.freeze_panes
    return worksheet


def _add_chart(worksheet, chart_spec, target: WorkbookSheet) -> None:
    rows = max((cell.row for cell in target.cells), default=1)
    columns = max((cell.column for cell in target.cells), default=1)
    _validate_chart_range(chart_spec.data_range, row_count=rows, column_count=columns)
    chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[chart_spec.type]()
    chart.title = chart_spec.title
    min_col, min_row, max_col, max_row = range_boundaries(chart_spec.data_range)
    chart.add_data(
        Reference(worksheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row),
        titles_from_data=chart_spec.titles_from_data,
    )
    if chart_spec.categories_range:
        cmin, rmin, cmax, rmax = range_boundaries(chart_spec.categories_range)
        _validate_chart_range(chart_spec.categories_range, row_count=rows, column_count=columns)
        chart.set_categories(
            Reference(worksheet, min_col=cmin, min_row=rmin, max_col=cmax, max_row=rmax)
        )
    worksheet.add_chart(chart, f"{get_column_letter(columns + 2)}2")


def _add_sources_sheet(workbook: Workbook, spec: WorkbookSpec) -> None:
    """Append a sheet legend of the sources behind generated content, under an unused name."""
    from app.artifacts.evidence import source_legend

    name, number = "Sources", 1
    while name.casefold() in {sheet.title.casefold() for sheet in workbook.worksheets}:
        number += 1
        name = f"Sources ({number})"
    legend = workbook.create_sheet(name)
    legend.append(["Element", "Source", "Locator", "Chunk ID"])
    for row in source_legend(spec):
        legend.append(list(row))
        for cell in legend[legend.max_row]:
            cell.data_type = "s"
    legend.freeze_panes = "A2"
    for column, width in {"A": 30, "B": 45, "C": 30, "D": 40}.items():
        legend.column_dimensions[column].width = width


def export_xlsx(spec: WorkbookSpec) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    worksheets = {sheet.id: _write_sheet(workbook, sheet) for sheet in spec.sheets}
    sheets = {sheet.id: sheet for sheet in spec.sheets}
    for chart_spec in spec.charts:
        _add_chart(worksheets[chart_spec.sheet_id], chart_spec, sheets[chart_spec.sheet_id])
    if spec.evidence:
        _add_sources_sheet(workbook, spec)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _unverified(xlsx_bytes: bytes, message: str) -> RecalculationResult:
    return RecalculationResult("needs_recalculation", xlsx_bytes, message)


def _recalculation_problems(xlsx_bytes: bytes, recalculated: bytes, spec: WorkbookSpec):
    """(changed, missing) cell addresses: inputs or formulas the runner altered, and formulas without a value."""
    # Runner output is untrusted too: apply the same archive preflight as uploads.
    inspect_office_archive(recalculated, kind="xlsx")
    formulas = load_workbook(io.BytesIO(recalculated), data_only=False, read_only=True)
    values = load_workbook(io.BytesIO(recalculated), data_only=True, read_only=True)
    original = load_workbook(io.BytesIO(xlsx_bytes), data_only=False, read_only=True)
    missing = []
    changed = []
    for sheet_spec in spec.sheets:
        form_ws = formulas[sheet_spec.name]
        value_ws = values[sheet_spec.name]
        for cell in sheet_spec.cells:
            address = sheet_spec.name + "!" + get_column_letter(cell.column) + str(cell.row)
            returned = form_ws.cell(cell.row, cell.column)
            if cell.formula is not None and returned.value != cell.formula:
                changed.append(address)
            if cell.formula is None:
                expected = original[sheet_spec.name].cell(cell.row, cell.column)
                if returned.value != expected.value or returned.data_type != expected.data_type:
                    changed.append(address)
            if cell.formula and (
                returned.data_type != "f"
                or value_ws.cell(cell.row, cell.column).value is None
                or value_ws.cell(cell.row, cell.column).data_type == "e"
            ):
                missing.append(address)
    formulas.close()
    values.close()
    original.close()
    return changed, missing


def recalculate_xlsx(
    xlsx_bytes: bytes,
    spec: WorkbookSpec,
    *,
    runner: IsolatedLibreOfficeRunner | None = None,
    timeout_seconds: int = 60,
) -> RecalculationResult:
    has_formulas = any(cell.formula is not None for sheet in spec.sheets for cell in sheet.cells)
    if not has_formulas:
        return RecalculationResult("not_required", xlsx_bytes)
    if runner is None:
        return _unverified(xlsx_bytes, "LibreOffice recalculation is not configured.")
    if not 1 <= timeout_seconds <= 300:
        raise ValueError("timeout_seconds must be between 1 and 300")
    try:
        recalculated = runner.recalculate_xlsx(xlsx_bytes, timeout_seconds=timeout_seconds)
    except Exception as exc:  # noqa: BLE001 - a sandbox runner failure cannot establish calculation status
        return _unverified(xlsx_bytes, f"Isolated recalculation failed: {exc}")
    if not recalculated:
        return _unverified(xlsx_bytes, "Recalculation produced no workbook bytes.")
    try:
        changed, missing = _recalculation_problems(xlsx_bytes, recalculated, spec)
    except Exception as exc:  # noqa: BLE001 - untrusted runner output must remain unverified
        return _unverified(xlsx_bytes, f"Could not verify calculated values: {exc}")
    if changed:
        return _unverified(
            xlsx_bytes, "Recalculation changed formulas or inputs at: " + ", ".join(changed[:20]))
    if missing:
        return _unverified(
            xlsx_bytes, "Calculated values are missing for: " + ", ".join(missing[:20]))
    return RecalculationResult("verified", recalculated)


def formula_preview(spec: WorkbookSpec) -> dict:
    """Return explicit formula text and unknown calculated values; never evaluate formulas here."""
    return {
        "status": "needs_recalculation"
        if any(c.formula for s in spec.sheets for c in s.cells)
        else "ready",
        "sheets": [
            {
                "id": sheet.id,
                "name": sheet.name,
                "cells": [
                    {
                        "id": cell.id,
                        "row": cell.row,
                        "column": cell.column,
                        "value": cell.value,
                        "formula": cell.formula,
                        "calculatedValue": None if cell.formula else cell.value,
                    }
                    for cell in sheet.cells
                ],
            }
            for sheet in spec.sheets
        ],
    }
